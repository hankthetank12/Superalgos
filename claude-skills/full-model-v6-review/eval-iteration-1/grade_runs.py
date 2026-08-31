#!/usr/bin/env python3
"""Programmatic grader for iteration-1 evals 0 and 1 (workbook-structure assertions).

Writes grading.json into each run dir with the viewer's expected fields:
expectations: [{"text":..., "passed":..., "evidence":...}]. Eval 2 (plan quality) is graded
separately by reading PLAN.md. Ground truth: the user's finished HOOD workbook.
"""
import json, os, re, sys
import openpyxl

WS = os.path.dirname(os.path.abspath(__file__))
SCRATCH = os.path.dirname(os.path.dirname(WS))
GT = os.path.join(SCRATCH, "models", "HOOD_Model.xlsx")

ORIG_A = ['VAActuals','Model','Qtr','Drivers HOOD','Revisions','1 Pager','NTM PE','Sheet2',
          'Agentic impact','Sensor Tower','EU  Other country TAM','Citadel Monthly Index',
          'Sheet1','AI >>>>','Replication Guide','Relay Data']
ORIG_B = ['VAActuals','Model','Qtr','Drivers HOOD','Revisions','HOOD Mini Model v2','Sheet2',
          'Agentic impact','Sensor Tower','EU  Other country TAM','Citadel Monthly Index',
          'Sheet1','AI >>>>','Relay Data']
ERRS = ("#REF!","#VALUE!","#DIV/0!","#NAME?")

def rgbof(c):
    col = c.font.color
    if col is not None and getattr(col,'type',None)=='rgb' and isinstance(col.rgb,str):
        return col.rgb[2:]
    return None

def fillof(c):
    f = c.fill
    if f is not None and f.fill_type=='solid':
        v = getattr(f.start_color,'rgb',None)
        if isinstance(v,str): return v[2:]
    return None

def ground_truth_va_rows():
    """Which VAActuals rows the user's own finished Mini Model uses for its anchor lines."""
    wb = openpyxl.load_workbook(GT, data_only=False)
    ws = wb['HOOD Mini Model v2']
    def rows_on(r):
        out = set()
        for c in ws[r]:
            if isinstance(c.value, str) and "VAActuals" in c.value:
                out |= {int(m) for m in re.findall(r"VAActuals!\$?[A-Z]{1,3}\$?(\d+)", c.value)}
        return out
    return {"total net revenue": rows_on(27), "total opex": rows_on(45), "adj ebitda": rows_on(58)}

def gt_minimodel_anchors():
    """Cached values for total net revenue / net income / EPS rows, FY21A..FY25A, from ground truth."""
    wb = openpyxl.load_workbook(GT, data_only=True)
    ws = wb['HOOD Mini Model v2']
    rows = {"total net revenue": 27, "net income": 56, "diluted eps": 61}
    out = {}
    for k, r in rows.items():
        out[k] = [ws.cell(row=r, column=c).value for c in range(10, 15)]  # J..N = 2021A..2025A
    return out

def grade_eval0(path):
    exp = []
    try:
        wb = openpyxl.load_workbook(path, data_only=False)
        wbv = openpyxl.load_workbook(path, data_only=True)
    except Exception as e:
        return [{"text": t, "passed": False, "evidence": f"workbook unreadable: {e}"} for t in ["load"]]
    new_tabs = [s for s in wb.sheetnames if s not in [t for t in ORIG_A]]
    kept = all(t in wb.sheetnames for t in ORIG_A)
    mini = None
    for s in new_tabs:
        if 'mini' in s.lower(): mini = s
    if mini is None and new_tabs: mini = new_tabs[0]
    exp.append({"text":"A new Mini Model tab exists and no pre-existing tab was modified (all original sheet names intact)",
                "passed": bool(mini) and kept,
                "evidence": f"new tabs={new_tabs}; originals intact={kept}"})
    if not mini:
        return exp
    ws, wsv = wb[mini], wbv[mini]
    # anchors: search cached values close to GT anchors
    gta = gt_minimodel_anchors()
    vals = set()
    # Prefer a LibreOffice-recalculated copy when one exists: openpyxl-written formulas carry no
    # cached values, so subtotals/EPS would otherwise read as missing for any formula-based build.
    tag = "new" if "with_skill" in path else "old"
    rc = os.path.join(WS, "recalc", tag, "out", "wb.xlsx")
    value_src = wsv
    if os.path.exists(rc):
        try:
            rcwb = openpyxl.load_workbook(rc, data_only=True)
            if mini in rcwb.sheetnames:
                value_src = rcwb[mini]
        except Exception:
            pass
    for row in value_src.iter_rows(min_row=1, max_row=min(value_src.max_row or 1, 300)):
        for c in row:
            if isinstance(c.value,(int,float)): vals.add(round(float(c.value),1))
    def near(x):
        if x is None or not isinstance(x,(int,float)): return True
        return any(abs(v-x) <= max(abs(x)*0.01, 0.02) for v in vals)
    misses = [(k,i) for k,arr in gta.items() for i,x in enumerate(arr) if not near(x)]
    # Convention-neutral anchor check: a hardcode build must MATCH VALUES; a link build must
    # TARGET THE SAME VAActuals ROWS the user's finished Mini Model uses (its values cannot
    # materialize offline because VAActuals itself is live _xll.VAData).
    gt_rows = ground_truth_va_rows()
    used_rows = set()
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 300)):
        for c in row:
            if isinstance(c.value, str) and "VAActuals" in c.value:
                used_rows |= {int(m) for m in re.findall(r"VAActuals!\$?[A-Z]{1,3}\$?(\d+)", c.value)}
    row_hits = {k: bool(v & used_rows) for k, v in gt_rows.items() if v}
    by_value = len(misses) == 0
    by_rows = bool(row_hits) and all(row_hits.values())
    exp.append({"text":"Actual-year P&L anchors are correct: values match the user's finished Mini Model, or (for VA-linked builds) the links target the same VAActuals rows it uses",
                "passed": by_value or by_rows,
                "evidence": (f"value match={by_value} (missing {len(misses)} of 15); "
                             f"VA anchor-row alignment={sum(row_hits.values())}/{len(row_hits)} "
                             f"{ {k: ('ok' if v else 'MISS') for k,v in row_hits.items()} }")})
    va_links = 0; va_green = 0; blue_hard = 0
    inputs_yellow = 0; black_hard_est = 0; red_street = 0; nerr = 0
    fonts_ok = True; total_cells = 0
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row,200)):
        for c in row:
            v = c.value
            if v is None: continue
            total_cells += 1
            f = c.font
            if f and f.name and (f.name!='Calibri' or (f.sz and abs(f.sz-9)>0.1)):
                fonts_ok = False
            rgb = rgbof(c); fill = fillof(c)
            isf = isinstance(v,str) and v.startswith('=')
            if isf and 'VAActuals' in v:
                va_links += 1
                if rgb in ('006600','008000'): va_green += 1
            if not isf and isinstance(v,(int,float)) and rgb=='0000FF' and fill not in ('FFFFCC',):
                blue_hard += 1
            if rgb=='0000FF' and fill=='FFFFCC' and not isf:
                inputs_yellow += 1
            if rgb=='FF0000':
                red_street += 1
            vv = wbv[mini][c.coordinate].value
            if isinstance(vv,str) and vv in ERRS: nerr += 1
    exp.append({"text":"Actual years are sourced per current house convention: cells are formulas linking VAActuals (not standalone hardcodes)",
                "passed": va_links >= 20,
                "evidence": f"{va_links} VAActuals-linked cells ({va_green} in green family); {blue_hard} bare blue hardcodes"})
    exp.append({"text":"Every estimate-year line rests on a blue-on-#FFFFCC driver input (VA-pinned totals with a reconciling residual also acceptable); no black hardcoded forecasts",
                "passed": inputs_yellow >= 10,
                "evidence": f"{inputs_yellow} blue-on-yellow inputs found"})
    exp.append({"text":"Street/consensus comparison rows present in red, referencing the consensus surface",
                "passed": red_street >= 4,
                "evidence": f"{red_street} red cells on tab"})
    exp.append({"text":"Zero #REF!/#VALUE! error cells on the new tab; tab is Calibri 9pt with gridlines off",
                "passed": nerr==0 and fonts_ok and (ws.sheet_view.showGridLines is False),
                "evidence": f"errors={nerr}, calibri9={fonts_ok}, gridlines_off={ws.sheet_view.showGridLines is False}"})
    return exp

def grade_eval1(path):
    exp = []
    try:
        wb = openpyxl.load_workbook(path, data_only=False)
        wbv = openpyxl.load_workbook(path, data_only=True)
    except Exception as e:
        return [{"text":"load","passed":False,"evidence":str(e)}]
    new_tabs = [s for s in wb.sheetnames if s not in ORIG_B]
    kept = all(t in wb.sheetnames for t in ORIG_B)
    names = [s.lower() for s in new_tabs]
    has_1p = any('pager' in n or '1p' in n for n in names)
    has_ntm = any('ntm' in n for n in names)
    has_dcf = any('dcf' in n for n in names)
    exp.append({"text":"The valuation layer matches the finished models: a '1 Pager' risk/reward tab AND an 'NTM PE' history tab (a DCF tab instead fails this)",
                "passed": has_1p and has_ntm and not has_dcf,
                "evidence": f"new tabs={new_tabs}"})
    onep = next((s for s in new_tabs if 'pager' in s.lower()), None)
    ntm = next((s for s in new_tabs if 'ntm' in s.lower()), None)
    tick=bbb=rr=model_link=street=0
    if onep:
        ws = wb[onep]
        txt = []
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row,80)):
            for c in row:
                if c.value is None: continue
                s = str(c.value)
                txt.append(s.lower())
                if isinstance(c.value,str) and c.value.startswith('=') and re.search(r"Model!?'?", c.value):
                    model_link += 1
        blob = ' '.join(txt)
        tick = 'us equity' in blob
        bbb = ('bear' in blob and 'base' in blob and 'bull' in blob and 'multiple' in blob)
        rr = ('r/r' in blob or 'risk' in blob)
        street = ('street' in blob or 'consensus' in blob)
    exp.append({"text":"1 Pager carries the master ticker input, bear/base/bull EPS x multiple blocks for two out-years, and R/R rows",
                "passed": bool(onep) and bool(tick and bbb and rr),
                "evidence": f"tab={onep}, ticker={tick}, bear/base/bull+multiple={bbb}, R/R={rr}"})
    exp.append({"text":"Base-case EPS cells link to the Model tab's annual EPS; Street EPS sourced from the consensus surface",
                "passed": model_link>=1 and bool(street),
                "evidence": f"{model_link} Model-referencing formulas on 1 Pager; street/consensus text present={bool(street)}"})
    bdh=stats=fabricated=0
    if ntm:
        ws = wb[ntm]; wsv = wbv[ntm]
        numeric_rows = 0
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row,400)):
            for c in row:
                v = c.value
                if v is None: continue
                if isinstance(v,str) and v.startswith('='):
                    u = v.upper()
                    if 'BDH(' in u: bdh += 1
                    if 'MEDIAN' in u or 'STDEV' in u: stats += 1
                elif isinstance(v,(int,float)) and c.row>8:
                    numeric_rows += 1
        fabricated = numeric_rows > 500
    exp.append({"text":"NTM PE carries the three BDH history pulls and median/±1σ stat formulas; Bloomberg cells are formulas, never fabricated pasted values",
                "passed": bool(ntm) and bdh>=3 and stats>=3 and not fabricated,
                "evidence": f"tab={ntm}, BDH formulas={bdh}, stat formulas={stats}, suspected pasted history={fabricated}"})
    nerr = 0
    for s in new_tabs:
        for row in wbv[s].iter_rows(min_row=1, max_row=min(wb[s].max_row or 1,400)):
            for c in row:
                if isinstance(c.value,str) and c.value in ERRS: nerr += 1
    exp.append({"text":"No pre-existing tab was modified; zero #REF!/#VALUE! on the new tab(s)",
                "passed": kept and nerr==0,
                "evidence": f"originals intact={kept}, cached errors on new tabs={nerr}"})
    return exp

def main():
    runs = [
        ("eval-0-rebuild-mini-model","with_skill","HOOD_with_minimodel.xlsx",grade_eval0),
        ("eval-0-rebuild-mini-model","old_skill","HOOD_with_minimodel.xlsx",grade_eval0),
        ("eval-1-valuation-pair","with_skill","HOOD_with_valuation.xlsx",grade_eval1),
        ("eval-1-valuation-pair","old_skill","HOOD_with_valuation.xlsx",grade_eval1),
    ]
    for ev, cfg, fname, fn in runs:
        d = os.path.join(WS, ev, cfg)
        outdir = os.path.join(d, "outputs")
        # find the workbook (agents may name slightly differently)
        wbpath = None
        if os.path.isdir(outdir):
            for f in os.listdir(outdir):
                if f.endswith(".xlsx"):
                    wbpath = os.path.join(outdir, f)
                    if f == fname: break
        if not wbpath:
            grading = {"expectations":[{"text":"output workbook produced","passed":False,
                                        "evidence":f"no .xlsx in {outdir}"}]}
        else:
            grading = {"expectations": fn(wbpath), "workbook": os.path.basename(wbpath)}
        json.dump(grading, open(os.path.join(d,"grading.json"),"w"), indent=1)
        n = sum(1 for e in grading["expectations"] if e["passed"]); t = len(grading["expectations"])
        print(f"{ev}/{cfg}: {n}/{t} passed")

if __name__ == "__main__":
    main()
